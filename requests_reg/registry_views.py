"""
API реестра доверенностей и МЧД (раздел юристов).

Доступ — юридическая служба (право legal_manage): видит и правит весь реестр.
Сотрудник со сквозным просмотром (системный администратор) — только читает.
Журнал правок — суперпользователю: право view_all или суперпользователь Django.
"""

import io
from datetime import timedelta

from django.db.models import Q
from django.http import FileResponse
from django.utils import timezone
from django.utils.decorators import method_decorator
from django.views.decorators.csrf import csrf_exempt
from rest_framework import serializers, viewsets
from rest_framework.decorators import action
from rest_framework.exceptions import AuthenticationFailed, NotFound, PermissionDenied
from rest_framework.permissions import AllowAny
from rest_framework.response import Response

from core.auth import can_view_all, get_current_b24_id, is_admin_mode, is_lawyer
from core.models import Organization, UserProfile
from core.search import query_param as _query_param

from . import constants as C
from . import registry
from .models import PoaNumbering, PoaRegistryEntry


class PoaRegistryEntrySerializer(serializers.ModelSerializer):
    organization_name = serializers.CharField(source="organization.short_name", read_only=True)
    kind_display = serializers.CharField(source="get_kind_display", read_only=True)
    request_number = serializers.SerializerMethodField()
    request_status = serializers.SerializerMethodField()
    request_status_display = serializers.SerializerMethodField()
    validity = serializers.SerializerMethodField()
    validity_display = serializers.SerializerMethodField()

    class Meta:
        model = PoaRegistryEntry
        fields = [
            "id", "request", "request_number", "request_status", "request_status_display",
            "kind", "kind_display", "organization", "organization_name",
            "number", "reg_number", "issue_date", "code", "representative", "powers",
            "valid_until", "cert_valid_until", "revoked_at", "comment",
            "validity", "validity_display", "created_at", "updated_at",
        ]
        read_only_fields = ["request", "created_at", "updated_at"]

    def get_request_number(self, obj):
        return obj.request.number if obj.request_id else ""

    def get_request_status(self, obj):
        return obj.request.status if obj.request_id else ""

    def get_request_status_display(self, obj):
        # статус доверенности в Мини-СЭД (ТЗ: в работе, на подписании, исполнено…)
        return obj.request.status_label() if obj.request_id else "Внесена вручную"

    def get_validity(self, obj):
        return registry.validity(obj, self.context.get("today"))

    def get_validity_display(self, obj):
        return dict(C.VALIDITY_CHOICES)[self.get_validity(obj)]


class PoaRegistryWriteSerializer(serializers.ModelSerializer):
    organization = serializers.PrimaryKeyRelatedField(queryset=Organization.objects.all())

    class Meta:
        model = PoaRegistryEntry
        fields = registry.EDITABLE_FIELDS


def validity_q(code: str, today) -> Q:
    """Отбор по состоянию срока — тем же правилом, что registry.validity."""
    revoked = Q(revoked_at__lte=today)
    live = ~revoked
    soon = today + timedelta(days=C.VALIDITY_WARN_DAYS)
    return {
        C.VALIDITY_REVOKED: revoked,
        C.VALIDITY_NO_TERM: live & Q(valid_until__isnull=True),
        C.VALIDITY_EXPIRED: live & Q(valid_until__lt=today),
        C.VALIDITY_EXPIRING: live & Q(valid_until__gte=today, valid_until__lte=soon),
        C.VALIDITY_ACTIVE: live & Q(valid_until__gt=soon),
    }.get(code, Q())


@method_decorator(csrf_exempt, name="dispatch")
class PoaRegistryViewSet(viewsets.ModelViewSet):
    permission_classes = [AllowAny]
    pagination_class = None
    serializer_class = PoaRegistryEntrySerializer

    def initial(self, request, *args, **kwargs):
        self.b24_id = get_current_b24_id(request)
        if not self.b24_id:
            raise AuthenticationFailed("Требуется авторизация Битрикс24.")
        super().initial(request, *args, **kwargs)
        # ТЗ: доступ к реестру — только юридическая служба (читать — ещё и
        # администратор со сквозным просмотром)
        if not (self._can_edit() or can_view_all(self.b24_id)):
            raise PermissionDenied("Реестр доступен только юридической службе.")

    def _can_edit(self) -> bool:
        return is_lawyer(self.b24_id) or is_admin_mode(self.request)

    def _require_edit(self):
        if not self._can_edit():
            raise PermissionDenied("Править реестр может только юридическая служба.")

    def _is_superuser(self) -> bool:
        user = getattr(self.request, "user", None)
        return can_view_all(self.b24_id) or bool(user and user.is_authenticated and user.is_superuser)

    def get_serializer_context(self):
        return {**super().get_serializer_context(), "today": timezone.localdate()}

    def get_queryset(self):
        qs = PoaRegistryEntry.objects.select_related("organization", "request")
        if self.action not in ("list", "export"):
            return qs
        p = self.request.query_params
        kind = p.get("kind")
        if kind in dict(C.REGISTRY_KINDS):
            qs = qs.filter(kind=kind)
        orgs = [x for x in (p.get("organization") or "").split(",") if x.isdigit()]
        if orgs:
            qs = qs.filter(organization_id__in=orgs)
        if p.get("validity"):
            qs = qs.filter(validity_q(p["validity"], timezone.localdate()))
        q = _query_param(self.request, "q").strip()
        if q:
            qs = qs.filter(
                Q(number__icontains=q) | Q(reg_number__icontains=q)
                | Q(representative__icontains=q) | Q(powers__icontains=q)
                | Q(code__icontains=q) | Q(comment__icontains=q)
                | Q(organization__short_name__icontains=q)
                | Q(request__number__icontains=q)
            )
        return qs

    # --- правка ---
    def _save(self, entry, *, partial=True):
        ser = PoaRegistryWriteSerializer(entry, data=self.request.data, partial=partial)
        ser.is_valid(raise_exception=True)
        return ser.validated_data

    def create(self, request, *args, **kwargs):
        """Ручная запись — доверенность, выданная до MiniSED или мимо него.
        assign_number=true — номер берётся из нумерации юрлица."""
        self._require_edit()
        ser = PoaRegistryWriteSerializer(data=request.data)
        ser.is_valid(raise_exception=True)
        data = dict(ser.validated_data)
        if request.data.get("assign_number") and not data.get("number"):
            # форма шлёт пустой «Номер» — его заменяет номер по маске
            number, seq, year = registry.allocate_number(
                data["organization"], data.get("issue_date") or timezone.localdate(),
            )
            data.update(number=number, seq=seq, seq_year=year)
        entry = PoaRegistryEntry.objects.create(**data, created_by_b24_id=self.b24_id)
        return Response(self.get_serializer(entry).data, status=201)

    def update(self, request, *args, **kwargs):
        self._require_edit()
        entry = self.get_object()
        data = self._save(entry, partial=kwargs.get("partial", False))
        registry.update_entry(entry, data, by_b24_id=self.b24_id)
        entry.refresh_from_db()
        return Response(self.get_serializer(entry).data)

    def destroy(self, request, *args, **kwargs):
        # Удалить можно только запись, внесённую руками: строка заявки — это
        # выданный номер, его не стирают (правят или помечают отзыв).
        self._require_edit()
        entry = self.get_object()
        if entry.request_id:
            raise PermissionDenied("Запись по заявке удалить нельзя — её можно только править.")
        return super().destroy(request, *args, **kwargs)

    # --- журнал правок (суперпользователь) ---
    @action(detail=True, methods=["get"])
    def history(self, request, pk=None):
        if not self._is_superuser():
            raise NotFound()
        entry = self.get_object()
        changes = list(entry.changes.all())
        ids = {c.changed_by_b24_id for c in changes if c.changed_by_b24_id}
        names = dict(UserProfile.objects.filter(bitrix_id__in=ids).values_list("bitrix_id", "fio"))
        labels = {f.name: str(f.verbose_name) for f in PoaRegistryEntry._meta.fields}
        return Response([
            {
                "id": c.id,
                "field": c.field,
                "field_display": labels.get(c.field, c.field),
                "old_value": c.old_value,
                "new_value": c.new_value,
                "changed_by_b24_id": c.changed_by_b24_id,
                "changed_by_name": names.get(c.changed_by_b24_id, ""),
                "changed_at": c.changed_at,
            }
            for c in changes
        ])

    # --- справочники экрана ---
    @action(detail=False, methods=["get"])
    def meta(self, request):
        orgs = Organization.objects.filter(is_active=True).order_by("short_name")
        return Response({
            "organizations": [{"id": o.id, "short_name": o.short_name} for o in orgs],
            "kinds": [{"code": c, "name": n} for c, n in C.REGISTRY_KINDS],
            "validity": [{"code": c, "name": n} for c, n in C.VALIDITY_CHOICES],
            "warn_days": C.VALIDITY_WARN_DAYS,
            "today": timezone.localdate(),
            "can_edit": self._can_edit(),
            "can_history": self._is_superuser(),
        })

    # --- нумерация: маски и счётчики юрлиц ---
    @action(detail=False, methods=["get", "patch"])
    def numbering(self, request):
        """GET — маски и последний номер по каждому юрлицу; PATCH
        {organization, mask?, last_seq?} — поправить (например, продолжить
        нумерацию, начатую в Excel)."""
        if request.method == "PATCH":
            self._require_edit()
            org = Organization.objects.filter(pk=request.data.get("organization")).first()
            if org is None:
                return Response({"detail": "Не выбрано юрлицо."}, status=400)
            n = registry.numbering_for(org)
            year = timezone.localdate().year
            if "mask" in request.data:
                mask = (request.data.get("mask") or "").strip()
                err = registry.mask_error(mask)
                if err:
                    return Response({"detail": err}, status=400)
                n.mask = mask
            if "last_seq" in request.data:
                try:
                    last = int(request.data.get("last_seq"))
                except (TypeError, ValueError):
                    return Response({"detail": "Последний номер — целое число."}, status=400)
                if last < 0:
                    return Response({"detail": "Последний номер не может быть отрицательным."}, status=400)
                n.year, n.last_seq = year, last
            n.save()
            from core.services import log_action

            log_action("poa_numbering_changed", target=n, request=request,
                       new_value={"mask": n.mask, "last_seq": n.last_seq, "by_b24_id": self.b24_id})
        today = timezone.localdate()
        out = []
        for org in Organization.objects.filter(is_active=True).order_by("short_name"):
            n = PoaNumbering.objects.filter(organization=org).first()
            mask = n.mask if n else registry.default_mask(org)
            last = n.last_seq if n and n.year == today.year else 0
            out.append({
                "organization": org.id,
                "organization_name": org.short_name,
                "mask": mask,
                "year": today.year,
                "last_seq": last,
                "next_number": registry.format_number(mask, last + 1, today),
            })
        return Response(out)

    # --- выгрузка в Excel ---
    @action(detail=False, methods=["get"])
    def export(self, request):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill

        today = timezone.localdate()
        wb = Workbook()
        ws = wb.active
        ws.title = "Реестр"
        ws.append(["Текущая дата:", today.strftime("%d.%m.%Y")])
        ws.append([])
        head = [
            "Компания", "Вид", "№", "Рег. № (ФНС)", "Дата выдачи", "КОД",
            "Представитель", "Полномочия", "Срок действия", "Срок сертификата",
            "Статус", "Статус в Мини-СЭД", "Заявка",
        ]
        ws.append(head)
        for cell in ws[3]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="5B9BD5")
            cell.alignment = Alignment(wrap_text=True, vertical="bottom")
        fills = {
            C.VALIDITY_ACTIVE: "A9D08E",
            C.VALIDITY_EXPIRING: "FFD966",
            C.VALIDITY_EXPIRED: "C00000",
            C.VALIDITY_REVOKED: "BFBFBF",
        }
        ser = PoaRegistryEntrySerializer(self.get_queryset(), many=True, context={"today": today})
        for row in ser.data:
            ws.append([
                row["organization_name"], row["kind_display"], row["number"], row["reg_number"],
                _d(row["issue_date"]), row["code"], row["representative"], row["powers"],
                _d(row["valid_until"]), _d(row["cert_valid_until"]),
                row["validity_display"], row["request_status_display"], row["request_number"],
            ])
            r = ws.max_row
            for cell in ws[r]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            for c in (5, 9, 10):
                ws.cell(r, c).number_format = "DD.MM.YYYY"
            color = fills.get(row["validity"])
            if color:
                ws.cell(r, 9).fill = PatternFill("solid", fgColor=color)
        for col, width in zip("ABCDEFGHIJKLM", (26, 13, 18, 38, 12, 16, 30, 70, 13, 13, 14, 20, 13)):
            ws.column_dimensions[col].width = width
        ws.auto_filter.ref = f"A3:M{max(ws.max_row, 3)}"
        ws.freeze_panes = "A4"
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return FileResponse(
            buf, as_attachment=True,
            filename=f"Реестр_доверенностей_{today:%Y-%m-%d}.xlsx",
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )


def _d(value):
    """ISO-дата из сериализатора → дата для ячейки Excel."""
    from datetime import date

    return date.fromisoformat(value) if value else None
